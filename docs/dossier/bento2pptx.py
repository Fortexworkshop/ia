"""Convertit le deck Bento en presentation PowerPoint native (Workshop2026-M1-G<n>-Pres.pptx).

Bento n'exporte pas en .pptx : il propose le PDF (impression) et le JSON du document. Ce
convertisseur relit le document du deck et reconstruit chaque element en **objet PowerPoint**
(zone de texte, forme, image, graphique), pour que le texte reste modifiable dans PowerPoint.

Correspondances :
  - 1 px du document (canvas 1280 x 720) = 9525 EMU = 0,75 pt  ->  les positions sont exactes ;
  - `text`  -> zone de texte, avec gras / italique / souligne et alignement ;
  - `shape` rect -> rectangle (ou rectangle arrondi) ; `shape` line -> connecteur a pointe ;
  - `image` -> capture WebP reencodee en PNG ou JPEG (PowerPoint ne lit pas le WebP) ;
  - `svg`   -> logo redessine en vectoriel (forme libre + ellipses), depuis `assets/logo.svg` ;
  - `chart` bar -> graphique a barres natif, categorie par categorie.

Usage :
    python docs/dossier/bento2pptx.py [--deck ../../rendus/neuf/SENTINEL-X-G20.bento.html]
                                      [--out ../../rendus/neuf]
"""

from __future__ import annotations

import argparse
import base64
import io
import json
import re
import subprocess
import sys
import tempfile
from html import unescape
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))  # docassets.py, voisin de ce script
import docassets  # noqa: E402  (apres l'ajout du dossier au chemin d'import)

ROOT = Path(__file__).resolve().parent.parent.parent      # Fortex/

PX = 9525                    # 1 px (96 dpi) en EMU
PT = 0.75                    # 1 px en points

FONT = "Segoe UI"            # pile systeme Windows, comme la charte (aucune police a installer)
MONO = "Consolas"

BLOCK = re.compile(r'<script type="application/bento\+json" id="bento-doc">(.*?)</script>', re.S)
TOKEN = re.compile(r'(<[^>]+>)')


# --------------------------------------------------------------------------- Texte enrichi

def runs(html: str):
    """Decoupe le html du deck en segments (texte, gras, italique, souligne)."""
    bold = italic = under = False
    out: list[tuple[str, bool, bool, bool]] = []
    for part in TOKEN.split(html):
        if not part:
            continue
        if part.startswith("<"):
            tag = part.strip("<>").split()[0].lower().rstrip("/")
            closing = part.startswith("</")
            if tag in ("b", "strong"):
                bold = not closing
            elif tag in ("i", "em"):
                italic = not closing
            elif tag == "u":
                under = not closing
            continue
        out.append((unescape(part), bold, italic, under))
    return out or [("", False, False, False)]


def lines_of(html: str):
    """Le <br> separe les paragraphes PowerPoint."""
    return [line for line in re.split(r"<br\s*/?>", html, flags=re.I)]


# --------------------------------------------------------------------------- Convertisseur

def convert(doc: dict, out_path: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt
    from PIL import Image

    W = Emu(doc["size"]["width"] * PX)
    H = Emu(doc["size"]["height"] * PX)

    def emu(value):
        return Emu(int(round(value * PX)))

    def rgb(value: str):
        value = (value or "").lstrip("#")
        if len(value) == 3:
            value = "".join(c * 2 for c in value)
        return RGBColor.from_string(value.upper())

    def opaque(value: str) -> bool:
        return bool(value) and value.lower() not in ("none", "transparent") \
            and not value.lower().startswith("rgba(0,0,0,0")

    anchors = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
    aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    tmp = Path(tempfile.mkdtemp(prefix="bento-pptx-"))
    pictures: dict[str, Path] = {}

    def picture_file(key: str) -> Path:
        """Les images du document sont des data URI : on les pose sur disque une seule fois.

        Les captures du deck sont en WebP ; PowerPoint ne lit pas ce format, on les reencode donc
        en JPEG (ou en PNG si l'image porte de la transparence).
        """
        if key in pictures:
            return pictures[key]
        raw = doc.get("assets", {}).get(key)
        if not raw or not raw.startswith("data:"):
            raise SystemExit(f"asset inconnu : {key}")
        image = Image.open(io.BytesIO(base64.b64decode(raw.split(",", 1)[1])))
        if image.mode in ("RGBA", "LA", "P"):
            path = tmp / f"{key}.png"
            image.convert("RGBA").save(path, "PNG", optimize=True)
        else:
            path = tmp / f"{key}.jpg"
            image.convert("RGB").save(path, "JPEG", quality=88, optimize=True)
        pictures[key] = path
        return path

    def add_text(slide, el):
        box = slide.shapes.add_textbox(emu(el["x"]), emu(el["y"]), emu(el["w"]), emu(el["h"]))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchors.get(el.get("valign", "top"), MSO_ANCHOR.TOP)
        line_height = el.get("lineHeight") or 1.3
        for index, line in enumerate(lines_of(el.get("html", ""))):
            para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
            para.alignment = aligns.get(el.get("align", "left"), PP_ALIGN.LEFT)
            para.line_spacing = line_height
            for content, bold, italic, under in runs(line):
                run = para.add_run()
                run.text = content
                font = run.font
                font.size = Pt(round(el.get("fontSize", 18) * PT, 1))
                font.bold = bold or el.get("fontWeight", 400) >= 600
                font.italic = italic
                font.underline = under
                font.name = MONO if "mono" in str(el.get("fontFamily", "")).lower() else FONT
                font.color.rgb = rgb(el.get("color", "#e6e9ee"))
                if el.get("letterSpacing"):
                    font._rPr.set("spc", str(int(el["letterSpacing"] * 100)))
        return box

    def add_rect(slide, el):
        radius = el.get("radius") or 0
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(shape_type, emu(el["x"]), emu(el["y"]), emu(el["w"]), emu(el["h"]))
        if radius > 0:
            shape.adjustments[0] = min(0.5, radius / max(1, min(el["w"], el["h"])))
        if opaque(el.get("fill", "")):
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(el["fill"])
        else:
            shape.fill.background()
        if el.get("strokeWidth") and opaque(el.get("stroke", "")):
            shape.line.color.rgb = rgb(el["stroke"])
            shape.line.width = Pt(max(0.75, el["strokeWidth"] * PT))
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        if el.get("strokeStyle") == "dashed":
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        shape.text_frame.text = ""
        return shape

    def add_line(slide, el):
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, emu(el["x"]), emu(el["y"]),
            emu(el["x"] + el["w"]), emu(el["y"] + el.get("h", 0)))
        connector.line.color.rgb = rgb(el.get("fill", "#79aaf7"))
        connector.line.width = Pt(max(0.75, el.get("strokeWidth", 2) * PT))
        if el.get("lineEnd") == "arrow":
            tail = connector.line._get_or_add_ln().makeelement(qn("a:tailEnd"), {"type": "triangle"})
            connector.line._get_or_add_ln().append(tail)
        return connector

    def add_image(slide, el, key=None):
        key = key or (el.get("src", "")[len("asset:"):] if el.get("src", "").startswith("asset:") else "")
        return slide.shapes.add_picture(str(picture_file(key)), emu(el["x"]), emu(el["y"]),
                                        emu(el["w"]), emu(el["h"]))

    def add_logo(slide, el):
        """Logo FORTEX en vectoriel : une forme libre pour le trace, deux ellipses pour le capteur.

        On le redessine plutot que de poser une image : `assets/logo.svg` reste la seule source de
        la forme, et le vecteur ne se pixellise ni a l'agrandissement ni a l'impression.
        """
        markup = el.get("markup", "")
        points, circles = docassets.logo_parts()
        stroke = re.search(r'<path\b[^>]*stroke="([^"]+)"', markup)
        accent = re.search(r'<circle\b[^>]*fill="([^"]+)"', markup)
        stroke = stroke.group(1) if stroke else docassets.LOGO_STROKE
        accent = accent.group(1) if accent else docassets.LOGO_ACCENT

        k = el["w"] / 32.0                 # le logo est dessine dans un repere 32 x 32
        ox, oy = el["x"], el["y"]

        def at(px, py):                    # point du repere du logo -> EMU de la diapositive
            return emu(ox + px * k), emu(oy + py * k)

        builder = slide.shapes.build_freeform(*at(*points[0]), scale=1)
        builder.add_line_segments([at(px, py) for px, py in points[1:]], close=True)
        trace = builder.convert_to_shape()
        trace.shadow.inherit = False
        trace.fill.background()
        trace.line.color.rgb = rgb(stroke)
        trace.line.width = emu(2 * k)

        for cx, cy, r, filled in circles:
            oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, *at(cx - r, cy - r),
                                          emu(2 * r * k), emu(2 * r * k))
            oval.shadow.inherit = False
            if filled:
                oval.fill.solid()
                oval.fill.fore_color.rgb = rgb(accent)
                oval.line.fill.background()
            else:
                oval.fill.background()
                oval.line.color.rgb = rgb(accent)
                oval.line.width = emu(1.5 * k)
                # PowerPoint n'a que des motifs de tirets imposes (dash = 4 x l'epaisseur) : on
                # repose le `stroke-dasharray="3 3"` du SVG en pourcentage de l'epaisseur du trait.
                line = oval.line._get_or_add_ln()
                custom = line.makeelement(qn("a:custDash"), {})
                custom.append(line.makeelement(qn("a:ds"), {"d": "200000", "sp": "200000"}))
                line.append(custom)

    def add_chart(slide, el):
        option = el.get("option", {})
        series = (option.get("series") or [{}])[0]
        data = CategoryChartData()
        data.categories = option.get("xAxis", {}).get("data", [])
        data.add_series(series.get("name", ""), tuple(series.get("data", [])))
        frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, emu(el["x"]), emu(el["y"]),
                                       emu(el["w"]), emu(el["h"]), data)
        chart = frame.chart
        chart.has_legend = False
        chart.has_title = False
        plot = chart.plots[0]
        plot.gap_width = 120
        plot.has_data_labels = False
        fill = series.get("itemStyle", {}).get("color", "#79aaf7")
        plot.series[0].format.fill.solid()
        plot.series[0].format.fill.fore_color.rgb = rgb(fill)
        axis_color = option.get("xAxis", {}).get("axisLabel", {}).get("color", "#a4adba")
        grid_color = option.get("yAxis", {}).get("splitLine", {}).get("lineStyle", {}).get("color", "#2a323e")
        for axis in (chart.category_axis, chart.value_axis):
            axis.has_major_gridlines = axis is chart.value_axis
            if axis.has_major_gridlines:
                axis.major_gridlines.format.line.color.rgb = rgb(grid_color)
                axis.major_gridlines.format.line.width = Pt(0.75)
            axis.format.line.color.rgb = rgb(grid_color)
            tick = axis.tick_labels.font
            tick.size = Pt(round(option.get("xAxis", {}).get("axisLabel", {}).get("fontSize", 13) * PT, 1))
            tick.color.rgb = rgb(axis_color)
            tick.name = FONT
        return chart

    for spec in doc["slides"]:
        slide = prs.slides.add_slide(blank)
        background = spec.get("background")
        if opaque(background or ""):
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = rgb(background)
        for el in spec.get("elements", []):
            kind = el.get("type")
            if kind == "shape":
                (add_line if el.get("shape") == "line" else add_rect)(slide, el)
            elif kind == "text":
                add_text(slide, el)
            elif kind == "image":
                add_image(slide, el)
            elif kind == "svg":
                add_logo(slide, el)
            elif kind == "chart":
                add_chart(slide, el)
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]

    prs.save(str(out_path))


# --------------------------------------------------------------------------- Entree

def load_doc(deck: Path) -> dict:
    source = deck.read_text(encoding="utf-8")
    match = BLOCK.search(source)
    if not match:
        raise SystemExit(f"bloc #bento-doc introuvable dans {deck}")
    return json.loads(match.group(1).replace("\\u003c", "<"))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--deck", default=str(ROOT.parent / "rendus/neuf/SENTINEL-X-G20.bento.html"),
                        help="deck Bento source")
    parser.add_argument("--out", default=str(ROOT.parent / "rendus/neuf"), help="dossier de sortie")
    parser.add_argument("--group", default="20", help="numero de groupe (<n> du nom de fichier)")
    args = parser.parse_args()

    deck = Path(args.deck)
    doc = load_doc(deck)
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    tag = f"G{args.group}-" if args.group else ""
    target = out_dir / f"Workshop2026-M1-{tag}Pres.pptx"
    convert(doc, target)
    print(f"PPTX : {target} ({target.stat().st_size // 1024} Ko, "
          f"{len(doc['slides'])} slides)")


if __name__ == "__main__":
    main()
